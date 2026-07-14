import io

from docx import Document as DocxDocument

from app.services.parsing.outline import HeadingHint


def extract_docx_pages(content: bytes) -> list[str]:
    pages, _outline = extract_docx_document(content)
    return pages


def extract_docx_document(content: bytes) -> tuple[list[str], list[HeadingHint]]:
    """Extract DOCX text and headings.

    python-docx has no concept of page boundaries, so the whole document is
    returned as a single page. Paragraph text plus table cell text are captured so
    spec tables (common in datasheets) are not dropped. Headings are detected from
    paragraph styles (``Heading 1/2/…`` or ``Title``) rather than font metrics.
    """
    document = DocxDocument(io.BytesIO(content))
    parts: list[str] = []
    headings: list[HeadingHint] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        parts.append(text)
        style_name = (paragraph.style.name if paragraph.style else "") or ""
        if style_name.startswith("Heading") or style_name == "Title":
            headings.append(HeadingHint(text=text, page=1, is_bold=True))
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return ["\n".join(parts)], headings
