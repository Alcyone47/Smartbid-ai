"""Excel export: a Summary sheet plus the full, styled Compliance Matrix."""

from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.reports.csv_report import _HEADER
from app.services.reports.data import ReportRow, status_label

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=15, bold=True)
_STATUS_FILL = {
    "match": PatternFill("solid", fgColor="C6EFCE"),
    "partial": PatternFill("solid", fgColor="FFEB9C"),
    "no_match": PatternFill("solid", fgColor="FFC7CE"),
}
_MATRIX_WIDTHS = [28, 44, 16, 11, 12, 10, 10, 18, 26, 12, 13, 50, 12]


def generate_matrix_xlsx(project: Any, rows: list[ReportRow], summaries: list) -> bytes:
    workbook = Workbook()
    _build_summary_sheet(workbook.active, project, summaries)
    _build_matrix_sheet(workbook.create_sheet("Compliance Matrix"), rows)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_summary_sheet(ws: Worksheet, project: Any, summaries: list) -> None:
    ws.title = "Summary"
    ws["A1"] = f"Compliance Summary — {project.name}"
    ws["A1"].font = _TITLE_FONT
    client = getattr(project, "client_name", None)
    if client:
        ws["A2"] = f"Client: {client}"
    ws["A3"] = f"Generated: {date.today().isoformat()}"

    header = ["Vendor", "Overall %", "Matched", "Partial", "Unmet", "Mandatory Met", "Optional Met"]
    header_row = 5
    for col, name in enumerate(header, start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for offset, s in enumerate(summaries, start=1):
        ws.cell(row=header_row + offset, column=1, value=s.vendor_name)
        ws.cell(row=header_row + offset, column=2, value=s.overall_compliance_pct)
        ws.cell(row=header_row + offset, column=3, value=s.matched)
        ws.cell(row=header_row + offset, column=4, value=s.partial)
        ws.cell(row=header_row + offset, column=5, value=s.unmatched)
        ws.cell(row=header_row + offset, column=6, value=f"{s.mandatory_met}/{s.mandatory_total}")
        ws.cell(row=header_row + offset, column=7, value=f"{s.optional_met}/{s.optional_total}")

    for col, width in enumerate([26, 11, 10, 10, 10, 15, 15], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _build_matrix_sheet(ws: Worksheet, rows: list[ReportRow]) -> None:
    for col, name in enumerate(_HEADER, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for r, row in enumerate(rows, start=2):
        values = [
            row.requirement_label,
            row.requirement_text,
            row.category or "",
            "Yes" if row.is_mandatory else "No",
            row.expected_value or "",
            row.unit or "",
            row.operator or "",
            row.vendor_name,
            row.vendor_value or "",
            status_label(row.status),
            row.compliance_pct,
            row.rationale,
            row.source_page if row.source_page is not None else "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value)
        status_cell = ws.cell(row=r, column=10)
        fill = _STATUS_FILL.get(row.status)
        if fill:
            status_cell.fill = fill

    for col, width in enumerate(_MATRIX_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    last_col = get_column_letter(len(_HEADER))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(rows) + 1)}"
    ws.freeze_panes = "A2"
