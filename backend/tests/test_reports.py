import io
from decimal import Decimal
from types import SimpleNamespace

import openpyxl

from app.services.matching.scoring import VendorComplianceSummary
from app.services.reports import (
    ReportRow,
    build_report_rows,
    generate_matrix_csv,
    generate_matrix_xlsx,
    generate_summary_pdf,
)
from tests._factories import make_requirement, make_specification


def _entry(status, score):
    return SimpleNamespace(status=status, match_score=Decimal(score), rationale="")


def _summary(vendor="Acme", pct=72.7):
    return VendorComplianceSummary(vendor, pct, 3, 2, 1, 0, 2, 2, 0, 0, 1, 0, 1, 1)


def _rows():
    return [
        ReportRow("core_switch", "Core Switch", "Throughput", ">= 20 Mbps", "Networking", True, "20", "Mbps", ">=", "Acme", "2 Gbps", "match", 100, "ok", 2),
        ReportRow("ups_unit", "UPS Unit", "UPS Capacity", "20 KVA", "Power", True, "20", "KVA", "=", "Acme", None, "no_match", 0, "not found", None),
    ]


def test_build_report_rows_matched_and_null_spec():
    matched = (
        _entry("match", "1.0"),
        make_requirement(key="throughput", label="Throughput", text=">= 20 Mbps", category="Networking"),
        make_specification(value="2", text="2 Gbps"),
        "Acme",
    )
    unmatched = (
        _entry("no_match", "0.0"),
        make_requirement(key="ups", label="UPS", text="20 KVA", is_mandatory=True),
        None,
        "Acme",
    )
    rows = build_report_rows([matched, unmatched])
    assert rows[0].vendor_value == "2"  # value preferred
    assert rows[0].category == "Networking"
    assert rows[0].compliance_pct == 100
    assert rows[1].vendor_value is None  # null spec
    assert rows[1].compliance_pct == 0


def test_csv_has_header_and_row_per_entry():
    content = generate_matrix_csv(_rows())
    text = content.decode("utf-8-sig")
    lines = [line for line in text.splitlines() if line.strip()]
    assert lines[0].startswith("Equipment,Requirement,Detail,Category")
    assert len(lines) == 3  # header + 2 rows
    assert "Throughput" in lines[1] and "UPS Capacity" in lines[2]


def test_xlsx_opens_with_both_sheets():
    project = SimpleNamespace(name="Enterprise SD-WAN", client_name="Globex")
    content = generate_matrix_xlsx(project, _rows(), [_summary()])
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Summary", "Compliance Matrix"]
    assert wb["Compliance Matrix"]["A1"].value == "Equipment"
    assert "Enterprise SD-WAN" in wb["Summary"]["A1"].value


def test_pdf_has_pdf_header():
    project = SimpleNamespace(name="Enterprise SD-WAN", client_name=None)
    content = generate_summary_pdf(project, [_summary()], _rows())
    assert content[:5] == b"%PDF-"
    assert len(content) > 800


def test_generators_handle_empty_rows():
    project = SimpleNamespace(name="Empty", client_name=None)
    assert generate_matrix_csv([]).decode("utf-8-sig").strip().startswith("Equipment")
    assert openpyxl.load_workbook(io.BytesIO(generate_matrix_xlsx(project, [], [])))
    assert generate_summary_pdf(project, [], [])[:5] == b"%PDF-"
